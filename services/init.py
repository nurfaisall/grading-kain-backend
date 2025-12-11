import pandas as pd
import os
from dotenv import load_dotenv
from openpyxl import load_workbook
import re

load_dotenv()

file_path = os.environ.get("DB_PATH")
db = [{"sheet": "db_zb", "table": "Table2"}, {"sheet": "db_pd", "table": "Table1"}]
data = []


def get_db():
    wb = load_workbook(file_path, read_only=False)
    for i in db:
        ws = wb[i["sheet"]]
        wb_zb = wb[i["sheet"]]
        wb_zb = wb_zb.tables[i["table"]].ref
        cols = re.sub(r"\d", "", wb_zb)
        wb_zb = re.sub(r"[a-zA-Z]", "", wb_zb).split(":")
        left = int(wb_zb[0])
        right = int(wb_zb[1])
        df = pd.read_excel(
            file_path,
            sheet_name=i["sheet"],
            usecols=cols,
            skiprows=left - 1,
            nrows=right - left + 1,
            header=0,
        )
        obj = {"name": i["sheet"], "data": df}
        data.append(obj)
    return
